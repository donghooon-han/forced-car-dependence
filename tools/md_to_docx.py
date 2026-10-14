"""
보고서 초안(.md)을 Word(.docx)로 바꾼다.
지원: # 제목, 문단, **굵게**, *기울임*, | 표 |, - 목록, 1. 번호 목록, > 인용, ---,
      '[그림 N 삽입: 경로]' 줄은 실제 그림으로 바꾼다.
사용: py tools/md_to_docx.py output/report_draft.md output/분석보고서_초안_v1.docx
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "맑은 고딕"
ROOT = Path(__file__).resolve().parents[1]


def set_font(run, size=None, bold=None, italic=None, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_inline(par, text, size=10, color=None, base_italic=False):
    """**굵게**, *기울임* 처리"""
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not tok:
            continue
        if tok.startswith("**"):
            set_font(par.add_run(tok[2:-2]), size, bold=True, italic=base_italic, color=color)
        elif tok.startswith("*"):
            set_font(par.add_run(tok[1:-1]), size, italic=True, color=color)
        else:
            set_font(par.add_run(tok), size, italic=base_italic, color=color)


def shade(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_table(doc, rows):
    header, body = rows[0], rows[2:]          # rows[1] 은 |---| 구분선
    t = doc.add_table(rows=1 + len(body), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(header):
        c = t.cell(0, j)
        c.text = ""
        add_inline(c.paragraphs[0], h, size=9)
        for r in c.paragraphs[0].runs:
            r.bold = True
        shade(c, "DCE6F0")
    for i, row in enumerate(body, start=1):
        for j in range(len(header)):
            c = t.cell(i, j)
            c.text = ""
            add_inline(c.paragraphs[0], row[j] if j < len(row) else "", size=9)
    doc.add_paragraph()


def hr(doc):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in {"w:val": "single", "w:sz": "6", "w:space": "1", "w:color": "999999"}.items():
        bottom.set(qn(k), v)
    b.append(bottom)
    pPr.append(b)


def convert(md_path: Path, out_path: Path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for side in ["left_margin", "right_margin"]:
        setattr(sec, side, Cm(2.0))
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(3)
    st.paragraph_format.line_spacing = 1.2
    width = sec.page_width - sec.left_margin - sec.right_margin

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        s = line.strip()
        if not s:
            i += 1
            continue
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            add_table(doc, rows)
            continue
        m = re.match(r"\[그림 (\d+) 삽입: (.+?)(?:\s*\|\s*([\d.]+))?\]", s)
        if m:
            img = ROOT / m.group(2).strip()
            ratio = float(m.group(3)) if m.group(3) else 1.0     # 본문 폭 대비 그림 폭
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.add_run().add_picture(str(img), width=int(width * ratio))
            i += 1
            continue
        if s == "---":
            hr(doc)
        elif s.startswith("#"):
            level = len(s) - len(s.lstrip("#"))
            text = s.lstrip("#").strip()
            size = {1: 15, 2: 13, 3: 11}.get(level, 10)
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10 if level <= 2 else 6)
            p.paragraph_format.keep_with_next = True
            add_inline(p, text.replace("**", ""), size=size, color="1F3864")
            for r in p.runs:
                r.bold = True
        elif s.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            add_inline(p, s.lstrip("> ").strip(), size=9, color="666666")
        elif re.match(r"^- ", s):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, s[2:])
        elif re.match(r"^\d+\. ", s):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\. ", "", s))
        elif s.startswith("*") and s.endswith("*") and not s.startswith("**"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, s[1:-1], size=9, color="444444", base_italic=True)
        else:
            p = doc.add_paragraph()
            add_inline(p, s)
        i += 1
    doc.save(out_path)
    print("저장:", out_path)


if __name__ == "__main__":
    convert(Path(sys.argv[1]), Path(sys.argv[2]))
